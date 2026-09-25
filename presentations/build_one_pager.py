from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P=Path(__file__).resolve().parent
r=Presentation();r.slide_width=Inches(16);r.slide_height=Inches(9);s=r.slides.add_slide(r.slide_layouts[6])
colors={'bg':'F5F7FA','ink':'172D3B','muted':'566875','red':'E51B2B','green':'157767','amber':'99641B','line':'DEE5EB','white':'FFFFFF','pink':'FCECEF'}
def color(c):return RGBColor.from_string(colors.get(c,c))
def box(x,y,w,h,c='white',border=None):
 a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=color(c)
 if border:a.line.color.rgb=color(border)
 else:a.line.fill.background()
 return a
def txt(t,x,y,w,h,size=18,c='ink',bold=False):
 a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=a.text_frame;f.word_wrap=True;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
 for i,v in enumerate(t.split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=v;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(c);p.space_after=Pt(0)
 return a
s.background.fill.solid();s.background.fill.fore_color.rgb=color('bg');box(0,0,.14,9,'red')
txt('FAN-AGENT',.55,.32,10,.55,32,bold=True)
txt('AI-assisted ceiling-fan CFD  |  Development status & scale-up plan',.57,.98,13.5,.42,20,'muted')
s.shapes.add_picture(str(P/'havells-logo.png'),Inches(14.45),Inches(.25),width=Inches(.95))
box(.55,1.62,14.9,.68,'ink');txt('Working prototype + live AI verified. Physical CFD validation is the release gate.',.8,1.8,14.3,.39,22,'white',True)
# Status panel
box(.55,2.6,6.1,3.65,'white','line');txt('WHERE WE ARE',.8,2.82,5.6,.4,21,'red',True)
status=[('Designer workflow','Havells UI, saved cases, STEP/STL import and checks.'),('Live Gemini planner','RPM proposals, missing-input handling and scope checks;\n3 live scenarios passed + 65 automated tests.'),('CFD development pipeline','Synthetic fan mesh and 2,000-iteration run completed;\nconvergence and wall treatment remain unqualified.'),('Reference evidence','25,920 public author air-speed values acquired;\nmatching blade geometry is still missing.')]
for i,(a,b) in enumerate(status):
 y=3.31+i*.72;txt(a,.82,y,5.5,.29,17,'ink',True);txt(b,.82,y+.29,5.55,.46,14,'muted')
# Milestones
box(6.92,2.6,8.53,3.65,'white','line');txt('NEXT STEPS / HOW TO SCALE',7.17,2.82,7.98,.4,21,'red',True)
rows=[('1','Persist comparison studies','Reviewed RPM cases, provenance and reload.','1–2 weeks'),('2','Qualify ceiling-fan CFD','Matched reference, mesh sensitivity and accuracy.','4–6 weeks*'),('3','Connect controlled execution','Uploaded CAD → mesh → solve → acceptance checks.','3–4 weeks*'),('4','Pilot with designers','Representative designs, usability and reporting.','2 weeks*')]
for i,(n,a,b,d) in enumerate(rows):
 y=3.4+i*.67;box(7.18,y,.32,.32,'red');txt(n,7.27,y+.025,.2,.24,13,'white',True);txt(a,7.65,y,5.3,.32,18,bold=True);txt(d,13.25,y,1.9,.34,17,'amber',True);txt(b,7.65,y+.32,7.35,.31,15,'muted')
# Dependencies strip
text_y=6.55;txt('REQUIREMENTS & DEPENDENCIES',.58,text_y,14.7,.35,20,'red',True)
blocks=[('DATA & VALIDATION','Matched blade CAD + measured airflow;\ntorque/power data for load validation.\nPublic evidence first; no CFD-team dependency.'),('DESIGN / TEST SUPPORT','Representative CAD and operating conditions\nfrom designers; pilot feedback. Test-lab\nmeasurements if public evidence is insufficient.'),('AI / IT & COMPUTE','Gemini API access/quota and data-policy review;\nWSL/OpenFOAM compute capacity; IT support\nfor deployment, access and security when scaling.')]
for i,(a,b) in enumerate(blocks):
 x=.55+i*5.04;box(x,7.07,4.82,1.16,'white','line');txt(a,x+.18,7.19,4.46,.25,15,'ink',True);txt(b,x+.18,7.51,4.46,.65,13,'muted')
txt('*Planning estimates, not commitments: assume one dedicated developer, suitable compute and required data ready.\nCFD qualification timing starts after matching geometry/data are available; execution and pilot gates depend on qualification.',.58,8.35,13.8,.48,12,'muted')
txt('23 SEP 2026',14.12,8.65,1.4,.22,10,'muted')
s.notes_slide.notes_text_frame.text='''Management one-pager, prepared 23 September 2026. All timing is an initial planning estimate, not an agreed delivery commitment. Milestone 1: 1–2 weeks. Milestone 2: 4–6 weeks after matched geometry/experimental data availability, potentially longer if numerical issues persist. Milestone 3: 3–4 weeks with interfaces stable; some implementation may overlap qualification but design-use execution remains gated. Milestone 4: 2 weeks after acceptable accuracy and a supported execution workflow. Do not add these durations into a calendar promise without staffing and data-readiness review.

No company CFD-team support is available and none is assumed. Public benchmark sourcing remains our responsibility. Designer-provided representative CAD/operating conditions and pilot feedback are future scale-up inputs; test-lab measurement access is a conditional requirement if adequate public validation evidence cannot be obtained. AI/IT support is for API quota, company data policy, compute and deployment, not supplying CFD expertise. AI cannot replace physics validation.

Evidence: fan-agent/docs/PROJECT_PLAN.md; docs/ai-live-checks/20260921T184638173355Z.json (3 live Gemini 2.5 Flash checks passed after a prompt correction); 65 local tests passed. Synthetic 168,591-cell ceiling-fan run completed 2,000 iterations but residual targets and wall treatment remain unqualified. Public author repositories provide 5,760 single-fan plus 20,160 two-fan speed values; these are not claimed checksum-identical to Dryad. Matched blade CAD is absent. Current capability is proposal/preparation and development diagnostics, not arbitrary uploaded-CAD execution or production accuracy.

All slide text, cards and panels are editable native PowerPoint objects. Only the supplied Havells logo is raster.'''
r.core_properties.title='Fan-Agent | Executive one-pager';r.core_properties.subject='Development status, estimated milestones and scaling dependencies'
for target in [P/'Fan-Agent_Executive_OnePager.pptx',Path('C:/Users/aksha/Downloads/Fan-Agent_Executive_OnePager.pptx')]:r.save(target)
print('Created editable one-slide deck.')
