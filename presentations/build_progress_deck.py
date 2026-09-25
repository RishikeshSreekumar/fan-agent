from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT=Path(__file__).resolve().parent
OUT=Path('C:/Users/aksha/Downloads/Fan-Agent_Project_Update_Editable.pptx')
LOGO=Path('C:/Users/aksha/AppData/Local/Temp/codex-clipboard-7c27c873-7f4b-415a-809a-3e5d59de22e7.png')
# Use the user-supplied mark; no report screenshot is embedded.
logo=Image.open(LOGO).convert('RGBA'); logo.save(ROOT/'havells-logo.png')
R=Presentation();R.slide_width=Inches(16);R.slide_height=Inches(9)
C={'red':'E51B2B','ink':'192D3A','muted':'586976','bg':'F5F7FA','line':'DFE5EA','green':'147968','amber':'A46919','pink':'FCECEF','white':'FFFFFF'}
def rgb(v): return RGBColor.from_string(C.get(v,v))
def box(s,x,y,w,h,fill='white',line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill);sh.line.fill.background() if line is None else None
 if line:sh.line.color.rgb=rgb(line)
 if radius:sh.adjustments[0]=.08
 return sh
def text(s,t,x,y,w,h,size=20,color='ink',bold=False):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(8)
 return sh
def base(title,sub,section):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('bg');box(s,0,0,.16,9,'red');text(s,section.upper(),.6,.32,12,.3,12,'red',True);text(s,title,.6,.86,13.3,.7,32,'ink',True);text(s,sub,.62,1.64,14,.6,16,'muted');s.shapes.add_picture(str(ROOT/'havells-logo.png'),Inches(14.48), Inches(.28),width=Inches(.92));box(s,.6,8.48,14.8,.015,'line');text(s,'HAVELLS  /  FAN-AGENT  •  INTERNAL PROJECT UPDATE',.6,8.65,10,.2,10,'muted');text(s,f'23 SEP 2026   |   {len(R.slides):02}',12.3,8.65,3.1,.2,10,'muted');return s
def card(s,title,body,x,y,w,h,accent='red',size=18):
 box(s,x,y,w,h,'white','line',True);box(s,x,y,.055,h,accent);text(s,title,x+.22,y+.18,w-.44,.6,21,accent,True);text(s,body,x+.22,y+.94,w-.44,h-1.03,size)
def notes(s,t):s.notes_slide.notes_text_frame.text=t

s=base('Fan-Agent','AI-assisted ceiling-fan design screening','Project briefing')
text(s,'From designer intent\nto controlled CFD studies',.75,2.58,9,1.55,38,'ink',True)
text(s,'A Havells workbench combining natural-language input,\nCAD preparation and deterministic engineering checks.',.78,4.5,9,1.1,23,'muted')
box(s,10.5,2.5,4.8,3.8,'ink',radius=True);text(s,'CURRENT POSITION',10.8,2.8,4.2,.4,14,'white',True);text(s,'Working foundation.\nLive AI verified.',10.8,3.43,4.1,1.1,29,'white',True);text(s,'CFD performance validation\nremains the critical gap.',10.8,4.97,4.1,.8,20,'white')
box(s,.78,6.65,14.5,1.1,'pink',radius=True);text(s,'Purpose',1,6.88,1.6,.35,19,'red',True);text(s,'Enable faster design iteration with explicit evidence and acceptance gates.',2.65,6.88,12,.55,22)
notes(s,'Status verified from source checkpoint and live Gemini evidence dated 21 September 2026. This is a development project, not a production-ready or physically validated CFD product. No delivery dates or quantified business savings have been committed.')

s=base('Project at a glance','A compact overview for management discussion','Executive summary')
labels=['Designer request','AI study proposal','Input & CAD checks','Controlled CFD','Engineering report']
for i,l in enumerate(labels):
 x=.65+i*3.02;box(s,x,2.48,2.8,.7,'ink' if i<3 else 'line',radius=True);text(s,l,x+.14,2.68,2.52,.3,16,'white' if i<3 else 'ink',True)
 if i<4:text(s,'›',x+2.83,2.63,.2,.4,23,'red',True)
text(s,'Implemented proposal / preparation flow',.68,3.27,8.5,.3,12,'green',True);text(s,'Execution and reporting integration pending',9.7,3.27,5.5,.3,12,'amber',True)
card(s,'What is available','Havells UI and saved cases\nSTEP/STL import and inspection\nGemini study proposals\nDevelopment solver diagnostics',.65,3.85,4.72,2.72,size=16)
card(s,'Why it matters','Less repetitive study preparation\nExplicit units and missing inputs\nTraceable engineering decisions\nA consistent designer workflow',5.64,3.85,4.72,2.72,size=16)
card(s,'What remains','Matched experimental geometry\nAccepted CFD convergence\nMesh and accuracy qualification\nUploaded-CAD execution workflow',10.63,3.85,4.72,2.72,'amber',16)
text(s,'Next product milestone: save an RPM comparison as one persistent study.',.8,7.25,14,.6,24,'ink',True)
notes(s,'The first three boxes show implemented capabilities, not a fully automated end-to-end pipeline. Benefits are intended outcomes and have not been quantified. Standardized reporting is a target; development diagnostics already exist.')

s=base('AI interprets. Engineering checks control.','A constrained workflow for ceiling-fan designers','System architecture')
card(s,'01  Designer input','Plain-language question\nRPM and installation details\nSTEP / STL supplied locally',.65,2.65,4.72,2.55,size=19)
card(s,'02  Gemini planner','Extract structured parameters\nIdentify missing information\nPropose up to four RPM cases',5.64,2.65,4.72,2.55,size=19)
card(s,'03  Deterministic checks','Validate units and dimensions\nReject unsupported settings\nKeep execution gated',10.63,2.65,4.72,2.55,'green',19)
box(s,.65,5.65,14.7,1.92,'ink',radius=True);text(s,'Planned execution path',.95,5.95,4.6,.4,22,'white',True);text(s,'Geometry → mesh → OpenFOAM → numerical checks → accepted metrics',.95,6.57,13.8,.5,24,'white')
notes(s,'Foam-Agent inspired the separation of interpretation, structured planning and review. Its planner and LLMService were inspected. No upstream multi-agent runtime is imported; no claim of operational upstream-agent reuse is made. Only typed request text is sent to the configured provider; uploaded CAD and saved cases are not included. The AI does not choose physics or launch solvers. OpenFOAM v2412 is the development runtime. MRF/SST remains an unqualified development recipe; AMI is later scope.')

s=base('Working capabilities, with evidence','Software checks and physical validation are tracked separately','Delivered to date')
rows=[('Designer workbench','Havells UI, case forms and saved records','Implemented'),('Geometry preparation','STEP/STL conversion, checks and preview','Implemented'),('Live AI planning','Complete, missing-input and unsupported requests','3 checks passed'),('Software regression','Input validation, transport and application tests','65 tests passed'),('Development CFD','Synthetic 3D ceiling-fan mesh and solver pipeline','2,000 iterations'),('Public reference data','Author datasets: single-fan and two-fan speeds','25,920 values')]
for i,(a,b,c) in enumerate(rows):
 y=2.48+i*.82;box(s,.65,y,14.7,.73,'white' if i%2==0 else 'EEF2F5',radius=True);text(s,a,.88,y+.18,3.4,.4,18,'ink',True);text(s,b,4.35,y+.18,7.4,.4,17);text(s,c,12.12,y+.18,3,.4,17,'green' if i<4 else 'amber',True)
notes(s,'Live AI evidence: docs/ai-live-checks/20260921T184638173355Z.json. First attempt incorrectly marked Compare unsupported; prompt corrected and three live checks passed. Six model calls across two rounds. 65 local tests and mocked browser smoke check passed. Synthetic mesh: 168,591 cells; execution completed but convergence acceptance failed. Dataset comprises 5,760 single-fan measurements and 20,160 two-fan values from pinned Berkeley author repositories. Single-fan coordinates normalized; two-fan data checked. Author CSV hashes differ from Dryad including after newline normalization. Do not claim exact Dryad equivalence.')

s=base('The critical gap is CFD confidence','A completed solver run is not yet a validated design prediction','Engineering readiness')
card(s,'Observed in development','Synthetic run completed normally\nShort-window flow / torque stabilized\nPressure residual target not met\nWall treatment remains unqualified',.65,2.65,7.1,3.45,'amber',21)
card(s,'Required before design use','Matched blade geometry + measurements\nConvergence or statistical stability\nMesh-sensitivity evidence\nQuantified error for each claimed metric',8.02,2.65,7.33,3.45,'red',21)
box(s,.65,6.55,14.7,1.13,'pink',radius=True);text(s,'Current boundary',.92,6.84,3,.4,20,'red',True);text(s,'No production accuracy claim. No automatic execution from uploaded CAD.',4.08,6.84,10.9,.5,21)
notes(s,'At 2,000 iterations pressure initial residual was 0.004486 against a provisional 0.001 target; velocity and k targets also unmet. Short-window torque variation 0.10%, last-snapshot airflow change 0.99%, longer-window torque drift 2.38%. y-plus approximately 2.83–604.75. Public Haiku dataset supplies air-speed data but no matched blade CAD. Speed magnitude is not axial velocity; electrical input power is not shaft power. Current development must not depend on unavailable company CFD-team assistance. FAN-01 ducted axial-fan research is paused and is not the ceiling-fan validation route.')

s=base('Next milestones and completion gates','Sequence driven by evidence, without uncommitted delivery dates','Delivery roadmap')
items=[('1','Persist comparison studies','Save reviewed RPM cases as one study; retain inputs and provenance.','Gate: reload a comparison without losing its case definitions.'),('2','Qualify ceiling-fan physics','Resolve reference geometry, numerical behavior and mesh sensitivity.','Gate: documented agreement with appropriate experimental evidence.'),('3','Connect controlled execution','Join CAD preparation, meshing, solving and bounded failure handling.','Gate: supported uploaded geometry completes an assessed workflow.'),('4','Deliver designer reports','Compare accepted airflow, torque, power and velocity distributions.','Gate: clear, reproducible outputs with visible acceptance status.')]
for i,(n,t,b,g) in enumerate(items):
 y=2.45+i*1.36;box(s,.66,y,.6,.6,'red',radius=True);text(s,n,.85,y+.12,.3,.3,20,'white',True);text(s,t,1.5,y,6.3,.42,22,'ink',True);text(s,b,1.5,y+.52,6.3,.66,17,'muted');box(s,8.35,y,7,1.08,'white','line',True);text(s,g,8.6,y+.22,6.5,.65,18,'green' if i==0 else 'amber')
notes(s,'Milestone 1 is the next agreed product task; it has not yet been implemented. Engineering qualification can progress separately, but execution and performance claims remain gated by it. No automatic 100% success promise for arbitrary CAD. Dates and staffing estimates should be agreed after the engineering evidence gap is resolved.')

s=base('A credible foundation for the next phase','What management can take away today','Meeting takeaway')
card(s,'Built','A usable ceiling-fan workbench\nCAD preparation and saved cases\nReal Gemini proposal generation\nA demonstrated CFD pipeline',.65,2.6,4.72,3.35,'green',21)
card(s,'Not yet established','Reliable production predictions\nA qualified operating envelope\nEnd-to-end uploaded-CAD runs\nValidated torque / airflow accuracy',5.64,2.6,4.72,3.35,'amber',21)
card(s,'Focus next','Persist reviewed comparisons\nClose the physical evidence gap\nIntegrate bounded execution\nValidate before broader rollout',10.63,2.6,4.72,3.35,'red',21)
box(s,.65,6.45,14.7,1.25,'ink',radius=True);text(s,'Progress is demonstrable. Engineering acceptance remains the release gate.',.97,6.78,14.05,.6,26,'white',True)
notes(s,'Presentation based on local project source and checkpoints through 21 September 2026; prepared 23 September 2026. Reference layout inspiration: user-provided Thermal_design_agent_brief 2.pdf. This deck does not inherit thermal-design capabilities or claims from that brief. All slide text, process boxes, cards and tables are native editable PowerPoint objects; only the user-supplied logo is raster. Source evidence paths: docs/PROJECT_PLAN.md, docs/solver-assessment.json, docs/convergence-study.md, docs/ai-study-planner.md, benchmarks/ceiling-fan-dryad/README.md. Public references: https://github.com/CenterForTheBuiltEnvironment/single-fan and https://github.com/CenterForTheBuiltEnvironment/two-fans .')
R.core_properties.title='Fan-Agent | Project progress and next milestones'
R.core_properties.subject='Havells ceiling-fan CFD workbench — management update'
R.core_properties.author='Fan-Agent Project'
R.save(OUT);R.save(ROOT/OUT.name)
# Structural check: no off-slide objects; content remains editable.
for i,s in enumerate(R.slides,1):
 for sh in s.shapes:
  assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=R.slide_width+100 and sh.top+sh.height<=R.slide_height+100,(i,sh.name)
print(str(OUT));print('Slides:',len(R.slides),'native text boxes:',sum(sh.has_text_frame for s in R.slides for sh in s.shapes))
